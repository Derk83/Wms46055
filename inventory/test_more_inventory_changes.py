import io
from unittest.mock import patch

import openpyxl
from django.contrib.auth.models import Permission, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from .models import InventoryItem


class InventoryMoreChangesTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser("admin-more", "admin@example.com", "pw")
        self.client.force_login(self.admin)
        self.item = InventoryItem.objects.create(
            part_number="PART-100",
            fb_part_number="FB-9001",
            model_number="MODEL-1",
            name="Fiber Bracket",
            description="Fiber mounting bracket",
            shipper="Do not show",
            building_room="Do not show",
            quantity_on_hand=7,
            rack="A",
            section="01",
            bin_location="02",
        )
        self.other = InventoryItem.objects.create(
            part_number="PART-200",
            fb_part_number="FB-9002",
            name="Other Item",
            quantity_on_hand=3,
        )

    def _xlsx_upload(self, rows):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in rows:
            sheet.append(row)
        stream = io.BytesIO()
        workbook.save(stream)
        return SimpleUploadedFile(
            "inventory.xlsx",
            stream.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_inventory_page_has_fb_part_number_without_removed_columns(self):
        response = self.client.get(reverse("inventory_list"))

        self.assertContains(response, '<th class="fb-col">FB Part #</th>', html=False)
        self.assertContains(response, "FB-9001")
        self.assertNotContains(response, '<th class="shipper-col">Shipper</th>', html=False)
        self.assertNotContains(response, 'data-label="Shipper"', html=False)
        self.assertNotContains(response, 'data-label="Building/Room"', html=False)
        self.assertNotContains(response, '<th class="more-col">More</th>', html=False)
        self.assertContains(response, "inventory-action-group")

    def test_inventory_general_search_finds_fb_part_number(self):
        response = self.client.get(reverse("inventory_list"), {"q": "FB-9001"})

        self.assertContains(response, "PART-100")
        self.assertNotContains(response, "PART-200")

    def test_inventory_has_dedicated_fb_filter_and_sorting(self):
        response = self.client.get(
            reverse("inventory_list"),
            {"fb_part_number": "9001", "sort": "fb_part_desc"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'name="fb_part_number"', html=False)
        self.assertContains(response, 'value="9001"', html=False)
        self.assertContains(response, "PART-100")
        self.assertNotContains(response, "PART-200")
        self.assertTrue(response.context["filters_active"])

        ascending = self.client.get(reverse("inventory_list"), {"sort": "fb_part"})
        descending = self.client.get(reverse("inventory_list"), {"sort": "fb_part_desc"})
        self.assertEqual(
            list(ascending.context["items"].values_list("fb_part_number", flat=True)),
            ["FB-9001", "FB-9002"],
        )
        self.assertEqual(
            list(descending.context["items"].values_list("fb_part_number", flat=True)),
            ["FB-9002", "FB-9001"],
        )

    def test_inventory_table_has_sticky_headers_and_safe_clickable_rows(self):
        response = self.client.get(reverse("inventory_list"))
        html = response.content.decode()

        self.assertEqual(response.status_code, 200)
        item_url = reverse("item_detail", args=[self.item.pk])
        self.assertIn(f'data-item-url="{item_url}"', html)
        self.assertIn(".inventory-table thead th { position: sticky; top: 68px;", html)
        self.assertIn("event.target.closest('a, button, input, select, textarea, label, form')", html)
        self.assertIn("window.location.assign(this.dataset.itemUrl)", html)

    def test_inventory_form_renders_and_saves_fb_part_number(self):
        response = self.client.get(reverse("inventory_edit", args=[self.item.pk]))
        self.assertContains(response, 'name="fb_part_number"', html=False)

        data = {
            "part_number": self.item.part_number,
            "model_number": self.item.model_number,
            "fb_part_number": "FB-UPDATED",
            "name": self.item.name,
            "category": "",
            "description": self.item.description,
            "shipper": self.item.shipper,
            "quantity_on_hand": self.item.quantity_on_hand,
            "unit": "each",
            "building_room": self.item.building_room,
            "rack": "A",
            "section": "01",
            "bin_location": "02",
            "low_stock_threshold": 0,
            "barcode_value": "",
            "qr_code_value": "",
            "active": "on",
        }
        response = self.client.post(reverse("inventory_edit", args=[self.item.pk]), data)

        self.assertEqual(response.status_code, 302)
        self.item.refresh_from_db()
        self.assertEqual(self.item.fb_part_number, "FB-UPDATED")

    def test_inventory_spreadsheet_import_and_export_include_fb_part_number(self):
        upload = self._xlsx_upload([
            ["Part #", "FB Part #", "Name", "Qty"],
            ["IMPORTED-1", "FB-IMPORT", "Imported Item", 4],
        ])
        response = self.client.post(reverse("inventory_import_xlsx"), {"spreadsheet": upload})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(InventoryItem.objects.get(part_number="IMPORTED-1").fb_part_number, "FB-IMPORT")

        response = self.client.get(reverse("export_inventory_xlsx"))
        workbook = openpyxl.load_workbook(io.BytesIO(response.content), data_only=True)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        self.assertIn("FB Part #", headers)
        rows = {row[0]: dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)}
        self.assertEqual(rows["PART-100"]["FB Part #"], "FB-9001")

    def test_material_request_picker_and_item_api_search_fb_part_number(self):
        response = self.client.get(reverse("material_request_create"))
        self.assertContains(response, "FB-9001")
        self.assertContains(response, "FB Part #")
        self.assertContains(response, "fb-9001")

        response = self.client.get(reverse("item_search_api"), {"q": "FB-9001"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["results"][0]["id"], self.item.pk)
        self.assertEqual(response.json()["results"][0]["fb_part_number"], "FB-9001")

    @patch("inventory.views.generate_qr_code", side_effect=lambda payload: payload)
    def test_location_qrs_encode_absolute_bbx_location_url(self, _generate):
        expected = "https://bbx.rplwms.com/locations/A-01-02/"

        response = self.client.get(reverse("location_list"))
        self.assertEqual(response.context["locations"][0]["location_url"], expected)
        self.assertEqual(response.context["locations"][0]["qr_b64"], expected)

        response = self.client.get(reverse("location_detail", args=["A-01-02"]))
        self.assertEqual(response.context["location_url"], expected)
        self.assertEqual(response.context["qr_b64"], expected)

    def test_scanners_accept_new_location_url_and_legacy_payload(self):
        absolute = "https://bbx.rplwms.com/locations/A-01-02/"
        api_response = self.client.get(reverse("scan_lookup_api"), {"code": absolute})
        self.assertEqual(api_response.status_code, 200)
        self.assertTrue(api_response.json()["is_location"])
        self.assertEqual(api_response.json()["location_key"], "A-01-02")

        browser_response = self.client.get(reverse("scan_lookup"), {"code": absolute})
        self.assertRedirects(browser_response, reverse("location_detail", args=["A-01-02"]))

        legacy_response = self.client.get(reverse("scan_lookup_api"), {"code": "LOC:A-01-02"})
        self.assertEqual(legacy_response.status_code, 200)


class ItemBarcodePermissionTests(TestCase):
    def setUp(self):
        self.item = InventoryItem.objects.create(part_number="UPC-1", name="UPC Item", quantity_on_hand=1)
        self.other = InventoryItem.objects.create(part_number="UPC-2", name="Other UPC Item", quantity_on_hand=1)
        self.procurement = self._user_with_perms(
            "procurement-more", "manage_item_barcodes", "view_inventoryitem"
        )
        self.manager = self._user_with_perms(
            "manager-more", "manage_item_barcodes", "view_inventoryitem"
        )
        self.logistics = self._user_with_perms(
            "logistics-more", "print_qr_codes", "scan_codes", "view_inventoryitem"
        )
        self.admin = User.objects.create_superuser("admin-barcode", "barcode@example.com", "pw")

    def _user_with_perms(self, username, *codenames):
        user = User.objects.create_user(username=username, password="pw")
        user.user_permissions.add(*Permission.objects.filter(codename__in=codenames))
        return user

    def test_allowed_roles_can_print_item_barcode_and_restricted_role_cannot(self):
        for user in (self.procurement, self.manager, self.admin):
            with self.subTest(user=user.username):
                self.client.force_login(user)
                response = self.client.get(reverse("item_barcode_print", args=[self.item.pk]))
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, "Print Barcode")

        self.client.force_login(self.logistics)
        response = self.client.get(reverse("item_barcode_print", args=[self.item.pk]))
        self.assertEqual(response.status_code, 302)

    def test_allowed_roles_can_assign_scanned_upc_and_leading_zeroes_are_preserved(self):
        for index, user in enumerate((self.procurement, self.manager, self.admin), start=1):
            self.client.force_login(user)
            value = f"00012345678{index}"
            response = self.client.post(
                reverse("item_upc_assign", args=[self.item.pk]),
                {"barcode_value": value},
            )
            self.assertRedirects(response, reverse("item_detail", args=[self.item.pk]))
            self.item.refresh_from_db()
            self.assertEqual(self.item.barcode_value, value)

    def test_restricted_role_cannot_assign_upc(self):
        self.client.force_login(self.logistics)
        response = self.client.post(
            reverse("item_upc_assign", args=[self.item.pk]),
            {"barcode_value": "012345678901"},
        )
        self.assertEqual(response.status_code, 302)
        self.item.refresh_from_db()
        self.assertIsNone(self.item.barcode_value)

    def test_duplicate_or_blank_upc_is_rejected(self):
        self.other.barcode_value = "012345678901"
        self.other.save()
        self.client.force_login(self.procurement)

        duplicate = self.client.post(
            reverse("item_upc_assign", args=[self.item.pk]),
            {"barcode_value": "012345678901"},
            follow=True,
        )
        self.assertContains(duplicate, "already assigned")
        blank = self.client.post(
            reverse("item_upc_assign", args=[self.item.pk]),
            {"barcode_value": ""},
            follow=True,
        )
        self.assertContains(blank, "required")
        self.item.refresh_from_db()
        self.assertIsNone(self.item.barcode_value)

    def test_item_detail_shows_camera_and_print_controls_only_to_allowed_roles(self):
        self.client.force_login(self.procurement)
        response = self.client.get(reverse("item_detail", args=[self.item.pk]))
        self.assertContains(response, reverse("item_barcode_print", args=[self.item.pk]))
        self.assertContains(response, reverse("item_upc_assign", args=[self.item.pk]))
        camera = self.client.get(reverse("item_upc_assign", args=[self.item.pk]))
        self.assertContains(camera, "html5-qrcode")

        self.client.force_login(self.logistics)
        response = self.client.get(reverse("item_detail", args=[self.item.pk]))
        self.assertNotContains(response, reverse("item_barcode_print", args=[self.item.pk]))
        self.assertNotContains(response, reverse("item_upc_assign", args=[self.item.pk]))

    @patch("inventory.models.generate_code128_barcode", return_value="encoded")
    def test_generated_barcode_image_uses_same_value_shown_on_label(self, generator):
        expected = self.item.auto_barcode_value
        self.assertEqual(self.item.barcode_image_b64, "encoded")
        generator.assert_called_once_with(expected)
