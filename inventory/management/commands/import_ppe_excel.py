from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

from inventory.models import InventoryItem, InventoryTransaction


CATEGORY_NAMES = {"SAFETY VESTS", "GLOVES", "HARD HATS", "SAFETY GLASSES"}


def make_part_number(category, name):
    prefix = "WH"
    if category:
        prefix = "".join(word[0] for word in category.split() if word)[:4].upper() or "WH"
    slug = "".join(ch if ch.isalnum() else "-" for ch in name.upper()).strip("-")
    while "--" in slug:
        slug = slug.replace("--", "-")
    return f"{prefix}-{slug[:40]}"


class Command(BaseCommand):
    help = "Import starter warehouse inventory from an Excel workbook."

    def add_arguments(self, parser):
        parser.add_argument("path", type=str)

    def handle(self, *args, **options):
        path = Path(options["path"])
        if not path.exists():
            raise CommandError(f"File not found: {path}")

        wb = load_workbook(path, data_only=True)
        ws = wb.active
        category = ""
        imported = 0
        skipped_notes = []

        for row in ws.iter_rows(values_only=True):
            name = row[0]
            qty = row[1]
            note = row[5] if len(row) > 5 else None
            if note:
                skipped_notes.append(str(note))
            if not name:
                continue
            name = str(name).strip()
            if name.upper() in CATEGORY_NAMES:
                category = name.upper()
                continue
            if name.upper() == "NAME":
                continue
            quantity = int(qty) if isinstance(qty, (int, float)) else 0
            part_number = make_part_number(category, name)
            barcode = part_number
            qr_value = part_number
            item, created = InventoryItem.objects.update_or_create(
                part_number=part_number,
                defaults={
                    "name": name,
                    "category": category,
                    "quantity_on_hand": quantity,
                    "barcode_value": barcode,
                    "qr_code_value": qr_value,
                    "building_room": "Warehouse",
                    "active": True,
                },
            )
            if created:
                InventoryTransaction.objects.create(
                    item=item,
                    transaction_type=InventoryTransaction.TransactionType.IMPORT,
                    quantity_delta=quantity,
                    notes=f"Initial import from {path.name}",
                )
            imported += 1

        self.stdout.write(self.style.SUCCESS(f"Imported/updated {imported} inventory items."))
        if skipped_notes:
            self.stdout.write("Side notes captured for manual review:")
            for note in skipped_notes:
                self.stdout.write(f"- {note}")
