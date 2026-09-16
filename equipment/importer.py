import hashlib
import json
import re
import zipfile
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from .models import (
    Asset,
    AssetComponent,
    AssetEvent,
    AssetIdentifier,
    ActiveCustody,
    Checkout,
    CheckoutItem,
    EquipmentCategory,
    EquipmentImportBatch,
    EquipmentImportRow,
    EquipmentLocation,
    EquipmentParty,
    EquipmentVendor,
    MaintenanceWorkOrder,
    RentalAsset,
    RentalContract,
)
from .services import serialize_equipment_mutation

PLACEHOLDER_SERIALS = {"", "N/A", "NA", "NO SERIAL #", "NO SERIAL", "NONE", "TBD"}
REVIEW_ROWS = {
    "Vehicles": {22},
    "Rentals": {15, 16, 18, 19},
    "Radios": {13},
    "Laser Rulers": {3},
    "Testers&Splicers": {72, 205, 206, 237, 238},
    "IT-IPADs": {12, 13},
}
SHEET_CODES = {
    "Vehicles": "VEH",
    "Rentals": "RNT",
    "Scanners": "SCN",
    "Radios": "RAD",
    "Laser Rulers": "LAS",
    "Testers&Splicers": "TST",
    "IT-IPADs": "IPD",
}
MAX_WORKBOOK_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024
MAX_ZIP_ENTRIES = 200
MAX_SHEET_ROWS = 5000
MAX_SHEET_COLUMNS = 100


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _json_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def _raw_row(ws, row_number):
    return {
        ws.cell(1, col).value if ws.cell(1, col).value is not None else f"column_{col}": _json_value(ws.cell(row_number, col).value)
        for col in range(1, ws.max_column + 1)
        if ws.cell(row_number, col).value is not None
    }


def _has_values(ws, row_number):
    return any(ws.cell(row_number, col).value not in (None, "") for col in range(1, ws.max_column + 1))


def _date_value(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not value:
        return None
    text = _text(value)
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _decimal(value):
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value).replace("$", "").replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None


def _category(name, code, requires_serial=False):
    category, _ = EquipmentCategory.objects.get_or_create(
        code=code,
        defaults={"name": name, "requires_serial": requires_serial},
    )
    return category


def _location(name):
    value = _text(name)
    if not value:
        return None
    code = re.sub(r"[^A-Z0-9]+", "-", value.upper()).strip("-")[:40] or "UNKNOWN"
    location, _ = EquipmentLocation.objects.get_or_create(code=code, defaults={"name": value})
    return location


def _party(name, department=""):
    value = _text(name)
    if not value or value.upper() in {"N/A", "NA", "NONE", "TBD"}:
        return None
    department = _text(department)
    kind = EquipmentParty.Kind.PERSON
    if value.upper() in {"WAREHOUSE", "RPL2", "MN", "CONSTRUCTION", "FIBER", "FIELD"}:
        kind = EquipmentParty.Kind.TEAM
    party = EquipmentParty.objects.filter(display_name__iexact=value, department__iexact=department).first()
    if party:
        return party
    return EquipmentParty.objects.create(kind=kind, display_name=value, department=department)


def _ownership(vendor):
    value = _text(vendor).upper()
    if "OWNED" in value or value in {"BB", "BLACK BOX"}:
        return Asset.Ownership.OWNED
    if "LEASE" in value:
        return Asset.Ownership.LEASED
    if value:
        return Asset.Ownership.RENTED
    return Asset.Ownership.UNKNOWN


def _status(source, party=None):
    value = _text(source).upper()
    if value in {"OOS", "OUT OF SERVICE"}:
        return Asset.Status.OUT_OF_SERVICE
    if value == "IN TRANSIT":
        return Asset.Status.IN_TRANSIT
    if value in {"RETURNED", "AVAILABLE"}:
        return Asset.Status.AVAILABLE
    if value in {"IN USE", "ISSUED"}:
        return Asset.Status.CHECKED_OUT if party else Asset.Status.AVAILABLE
    if value == "ACTIVE" and party:
        return Asset.Status.CHECKED_OUT
    return Asset.Status.AVAILABLE if not value else Asset.Status.UNKNOWN


def _next_tag(sequence):
    return f"RPL-EQ-{sequence:06d}"


def _initial_sequence():
    highest = 0
    for tag in Asset.objects.filter(asset_tag__startswith="RPL-EQ-").values_list("asset_tag", flat=True):
        match = re.fullmatch(r"RPL-EQ-(\d+)", tag)
        if match:
            highest = max(highest, int(match.group(1)))
    return highest + 1


def _identifier(asset, kind, value, namespace, primary=False):
    value = _text(value)
    if value.upper() in PLACEHOLDER_SERIALS:
        return None
    normalized = " ".join(value.upper().split())
    if AssetIdentifier.objects.filter(
        namespace=namespace, kind=kind, normalized_value=normalized
    ).exists():
        note = f"Duplicate {kind.lower()} identifier in namespace {namespace}: {value}"
        asset.review_required = True
        asset.review_notes = "; ".join(filter(None, (asset.review_notes, note)))
        asset.save(update_fields=("review_required", "review_notes", "updated_at"))
        return None
    return AssetIdentifier.objects.create(
        asset=asset,
        kind=kind,
        namespace=namespace,
        value=value,
        is_primary=primary,
    )


def _create_asset(*, batch, actor, sequence, sheet, row, category, legacy_tag, name, manufacturer="", model="", ownership=Asset.Ownership.UNKNOWN, status=Asset.Status.AVAILABLE, party=None, location=None, acquired_on=None, cost=None, quantity=1, notes="", attributes=None, review=False, review_notes=""):
    asset = Asset.objects.create(
        asset_tag=_next_tag(sequence),
        legacy_tag=_text(legacy_tag),
        source_namespace=sheet,
        category=category,
        name=_text(name) or category.name,
        manufacturer=_text(manufacturer),
        model_number=_text(model),
        ownership=ownership,
        status=status,
        condition=Asset.Condition.UNKNOWN,
        quantity=max(int(quantity or 1), 1),
        current_party=party,
        home_location=location,
        current_location=location,
        acquired_on=acquired_on,
        purchase_cost=cost,
        notes=_text(notes),
        attributes=attributes or {},
        review_required=review,
        review_notes=review_notes,
        created_by=actor,
        updated_by=actor,
    )
    AssetEvent.objects.create(
        asset=asset,
        actor=actor,
        event_type=AssetEvent.Type.IMPORTED,
        summary=f"Imported from {sheet} row {row}",
        metadata={"batch": str(batch.pk), "sheet": sheet, "row": row},
    )
    if status == Asset.Status.CHECKED_OUT and party:
        checkout = Checkout.objects.create(
            checkout_number=f"EQ-CO-IMP-{sequence:06d}",
            borrower=party,
            borrower_snapshot=party.display_name,
            destination=location,
            checked_out_at=timezone.now(),
            purpose=f"Imported baseline custody from {sheet} row {row}",
            created_by=actor,
        )
        checkout_item = CheckoutItem.objects.create(
            checkout=checkout,
            asset=asset,
            asset_tag_snapshot=asset.asset_tag,
            asset_name_snapshot=asset.name,
            condition_out=asset.condition,
        )
        ActiveCustody.objects.create(
            asset=asset, checkout_item=checkout_item, borrower=party
        )
        AssetEvent.objects.create(
            asset=asset,
            actor=actor,
            event_type=AssetEvent.Type.CHECKED_OUT,
            summary=f"Baseline custody imported for {party.display_name}",
            metadata={"checkout": checkout.checkout_number, "source": "workbook_import"},
        )
    return asset


def _stage(batch, ws, row, status, normalized=None, messages=None, asset=None):
    messages = list(messages or [])
    if asset and status == EquipmentImportRow.Status.ACCEPTED and asset.review_required:
        status = EquipmentImportRow.Status.REVIEW
        if asset.review_notes:
            messages.append(asset.review_notes)
    return EquipmentImportRow.objects.create(
        batch=batch,
        sheet_name=ws.title,
        row_number=row,
        status=status,
        raw_data=_raw_row(ws, row),
        normalized_data=normalized or {},
        messages=messages,
        asset=asset,
    )


@transaction.atomic
def import_workbook(content, *, source_name, actor):
    if hasattr(content, "read"):
        content = content.read()
    if len(content) > MAX_WORKBOOK_BYTES:
        raise ValueError("Workbook exceeds the 10 MB upload limit.")
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise ValueError("Workbook contains too many internal files.")
            if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("Workbook expands beyond the 25 MB safety limit.")
            for entry in entries:
                if entry.compress_size and entry.file_size / entry.compress_size > 100:
                    raise ValueError("Workbook contains an unsafe compression ratio.")
    except zipfile.BadZipFile as error:
        raise ValueError("The upload is not a valid XLSX workbook.") from error
    digest = hashlib.sha256(content).hexdigest()
    workbook = load_workbook(BytesIO(content), data_only=True, read_only=False)
    for worksheet in workbook.worksheets:
        if worksheet.max_row > MAX_SHEET_ROWS or worksheet.max_column > MAX_SHEET_COLUMNS:
            raise ValueError(f"Worksheet {worksheet.title} exceeds safe row or column limits.")
    serialize_equipment_mutation()
    existing = EquipmentImportBatch.objects.filter(sha256=digest).first()
    if existing:
        return existing, False
    batch = EquipmentImportBatch.objects.create(source_name=source_name, sha256=digest, created_by=actor)
    sequence = _initial_sequence()
    accepted = review = excluded = total = 0
    created_assets = []

    def register(ws, row, asset, is_review=False, messages=None):
        nonlocal accepted, review, total
        if asset.category.requires_serial and not asset.identifiers.exists():
            asset.review_required = True
            asset.review_notes = "; ".join(filter(None, (asset.review_notes, "Required serial or external identifier is missing.")))
            asset.save(update_fields=("review_required", "review_notes", "updated_at"))
        is_review = is_review or asset.review_required
        state = EquipmentImportRow.Status.REVIEW if is_review else EquipmentImportRow.Status.ACCEPTED
        _stage(batch, ws, row, state, {"asset_tag": asset.asset_tag}, messages, asset)
        total += 1
        if is_review:
            review += 1
        else:
            accepted += 1
        created_assets.append(asset)

    # Vehicles
    ws = workbook["Vehicles"]
    category = _category("Vehicles", "VEHICLE", True)
    for row in range(2, 28):
        if not _has_values(ws, row):
            continue
        vendor, vin, year, make, plate, subtype, legacy, holder, user, team, start, source_status, cost, notes, gps_serial, gps_imei = [ws.cell(row, col).value for col in range(1, 17)]
        party = _party(user or holder, team)
        location = _location(team)
        bad_date = bool(start) and _date_value(start) is None
        needs_review = row in REVIEW_ROWS[ws.title] or bad_date
        asset = _create_asset(batch=batch, actor=actor, sequence=sequence, sheet=ws.title, row=row, category=category, legacy_tag=legacy, name=f"{_text(year)} {_text(make)} {_text(subtype)}".strip(), manufacturer=_text(make).split("(")[0].strip(), model=subtype, ownership=_ownership(vendor), status=_status(source_status, party), party=party, location=location, acquired_on=_date_value(start), cost=_decimal(cost), notes=notes, attributes={"rental_company": _text(vendor), "assigned_under": _text(holder), "year": _text(year)}, review=needs_review, review_notes="Invalid source date" if bad_date else "Source row requires reconciliation")
        sequence += 1
        _identifier(asset, AssetIdentifier.Kind.VIN, vin, "global", True)
        _identifier(asset, AssetIdentifier.Kind.LICENSE_PLATE, plate, "Vehicles")
        _identifier(asset, AssetIdentifier.Kind.GPS_SERIAL, _text(gps_serial).removeprefix("S/N:").strip(), "Vehicles")
        _identifier(asset, AssetIdentifier.Kind.IMEI, _text(gps_imei).removeprefix("IMEI:").strip(), "global")
        match = re.search(r"\((\d+)\)", _text(make))
        if match:
            _identifier(asset, AssetIdentifier.Kind.EQUIPMENT_NUMBER, match.group(1), "Vehicles")
        register(ws, row, asset, needs_review, ["Invalid date text preserved"] if bad_date else None)

    # Rentals: each row is a tracked rental line; quantity remains explicit.
    ws = workbook["Rentals"]
    category = _category("Rental Equipment", "RENTAL")
    contracts = {}
    for row in range(2, ws.max_row + 1):
        if not _has_values(ws, row):
            continue
        vendor, po, equipment_type, equipment_number, make, model, quantity, user, source_status, cost, start, return_date, contact = [ws.cell(row, col).value for col in range(1, 14)]
        if row > 20 or not equipment_type:
            _stage(batch, ws, row, EquipmentImportRow.Status.ARTIFACT, messages=["Non-record worksheet artifact preserved"])
            total += 1
            excluded += 1
            continue
        vendor_name = _text(vendor) or "Unknown rental vendor"
        vendor_obj, _ = EquipmentVendor.objects.get_or_create(name=vendor_name)
        contract = contracts.get(vendor_name)
        if not contract:
            contract = RentalContract.objects.create(contract_number=f"IMPORT-{batch.sha256[:8].upper()}-{len(contracts)+1}", vendor=vendor_obj, po_number=_text(po), starts_on=_date_value(start), status=RentalContract.Status.ACTIVE, created_by=actor, notes="Imported rental lines")
            contracts[vendor_name] = contract
        party = _party(user)
        needs_review = row in REVIEW_ROWS[ws.title]
        asset = _create_asset(batch=batch, actor=actor, sequence=sequence, sheet=ws.title, row=row, category=category, legacy_tag=equipment_number, name=equipment_type, manufacturer=make, model=model, ownership=Asset.Ownership.RENTED, status=_status(source_status, party), party=party, location=_location(user if _text(user).upper() in {"WAREHOUSE", "RPL2", "MN"} else ""), acquired_on=_date_value(start), cost=_decimal(cost), quantity=int(quantity or 1), notes=contact, attributes={"po_number": _text(po), "source_return_date": _json_value(return_date)}, review=needs_review, review_notes="Rental identifier or start date requires reconciliation" if needs_review else "")
        sequence += 1
        _identifier(asset, AssetIdentifier.Kind.EQUIPMENT_NUMBER, equipment_number, f"Rental:{vendor_name}", True)
        RentalAsset.objects.create(contract=contract, asset=asset, vendor_equipment_number=_text(equipment_number), rate_amount=_decimal(cost), expected_return_on=_date_value(return_date))
        register(ws, row, asset, needs_review)

    # Scanners
    ws = workbook["Scanners"]
    category = _category("Scanners", "SCANNER", True)
    for row in range(2, 53):
        if not _has_values(ws, row):
            continue
        legacy, user, department, model, unit_serial, base_serial, _, preferred = [ws.cell(row, col).value for col in range(1, 9)]
        party = _party(user, department)
        asset = _create_asset(batch=batch, actor=actor, sequence=sequence, sheet=ws.title, row=row, category=category, legacy_tag=legacy, name="Barcode scanner", manufacturer="Zebra", model=model, ownership=Asset.Ownership.OWNED, status=_status("In Use" if party else "Available", party), party=party, location=_location(department), attributes={"preferred_model_annotation": _text(preferred)})
        sequence += 1
        _identifier(asset, AssetIdentifier.Kind.SERIAL, unit_serial, "Scanners", True)
        _identifier(asset, AssetIdentifier.Kind.BASE_SERIAL, base_serial, "Scanners")
        AssetComponent.objects.create(asset=asset, role="Scanner unit", serial_number=_text(unit_serial), model_number=_text(model), source_row=row)
        AssetComponent.objects.create(asset=asset, role="Charging base", serial_number=_text(base_serial), source_row=row)
        register(ws, row, asset)

    # Radios
    ws = workbook["Radios"]
    category = _category("Radios", "RADIO", True)
    for row in range(2, 14):
        if not _has_values(ws, row):
            continue
        legacy, device_id, user, department, notes = [ws.cell(row, col).value for col in range(1, 6)]
        party = _party(user, department)
        needs_review = row in REVIEW_ROWS[ws.title]
        asset = _create_asset(batch=batch, actor=actor, sequence=sequence, sheet=ws.title, row=row, category=category, legacy_tag=legacy, name="Two-way radio", ownership=Asset.Ownership.OWNED, status=_status("In Use" if party else "Available", party), party=party, location=_location(department), notes=notes, review=needs_review, review_notes="Assignment requires confirmation" if needs_review else "")
        sequence += 1
        _identifier(asset, AssetIdentifier.Kind.SERIAL, device_id, "Radios", True)
        register(ws, row, asset, needs_review)

    # Laser rulers
    ws = workbook["Laser Rulers"]
    category = _category("Laser Rulers", "LASER", True)
    for row in range(2, 6):
        if not _has_values(ws, row):
            continue
        legacy, device_id, user, department, returned, notes = [ws.cell(row, col).value for col in range(1, 7)]
        party = _party(user, department)
        needs_review = row in REVIEW_ROWS[ws.title]
        asset = _create_asset(batch=batch, actor=actor, sequence=sequence, sheet=ws.title, row=row, category=category, legacy_tag=legacy, name="Laser ruler", ownership=Asset.Ownership.OWNED, status=_status("In Use" if party else "Available", party), party=party, location=_location(department), notes=notes, attributes={"source_return_date": _json_value(returned)}, review=needs_review, review_notes="Empty returned box / missing equipment requires investigation" if needs_review else "")
        sequence += 1
        _identifier(asset, AssetIdentifier.Kind.SERIAL, device_id, "Laser Rulers", True)
        register(ws, row, asset, needs_review)

    # Testers and splicers; continuation rows become components of the prior MMC.
    ws = workbook["Testers&Splicers"]
    categories = {}
    prior_asset = None
    for row in range(2, ws.max_row + 1):
        if not _has_values(ws, row):
            continue
        vendor, po, equipment_type, legacy, serial, make, model, quantity, user, source_status, cost, start, return_date, notes = [ws.cell(row, col).value for col in range(1, 15)]
        if not legacy and (serial or model) and prior_asset is not None:
            AssetComponent.objects.create(asset=prior_asset, role="MMC component", serial_number=_text(serial).strip("()"), model_number=_text(model), quantity=int(quantity or 1), notes=_text(notes), source_row=row)
            _stage(batch, ws, row, EquipmentImportRow.Status.ACCEPTED, {"component_of": prior_asset.asset_tag}, ["Continuation row imported as serialized component"], prior_asset)
            total += 1
            continue
        if not legacy:
            _stage(batch, ws, row, EquipmentImportRow.Status.ARTIFACT, messages=["Non-asset row preserved"])
            total += 1
            excluded += 1
            continue
        category_name = _text(equipment_type) or "Test Equipment"
        code = "TEST-" + re.sub(r"[^A-Z0-9]", "", category_name.upper())[:20]
        category = categories.setdefault(category_name, _category(category_name, code, True))
        party = _party(user)
        needs_review = row in REVIEW_ROWS[ws.title]
        placeholder = _text(serial).upper() in PLACEHOLDER_SERIALS
        asset = _create_asset(batch=batch, actor=actor, sequence=sequence, sheet=ws.title, row=row, category=category, legacy_tag=legacy, name=category_name, manufacturer=make, model=model, ownership=_ownership(vendor), status=_status(source_status, party), party=party, acquired_on=_date_value(start), cost=_decimal(cost), quantity=int(quantity or 1), notes=notes, attributes={"po_number": _text(po), "source_return_date": _json_value(return_date)}, review=needs_review, review_notes="Missing serial, quantity, or accessory exception requires review" if needs_review else "")
        sequence += 1
        _identifier(asset, AssetIdentifier.Kind.SERIAL, _text(serial).strip("()"), ws.title, True)
        if category_name == "MMC":
            AssetComponent.objects.create(asset=asset, role="MMC component", serial_number=_text(serial).strip("()"), model_number=_text(model), quantity=int(quantity or 1), source_row=row)
        if _ownership(vendor) == Asset.Ownership.RENTED:
            vendor_obj, _ = EquipmentVendor.objects.get_or_create(name=_text(vendor))
            key = f"TEST-{vendor_obj.pk}"
            contract = contracts.get(key)
            if not contract:
                contract = RentalContract.objects.create(contract_number=f"IMPORT-{batch.sha256[:8].upper()}-TEST", vendor=vendor_obj, po_number=_text(po), starts_on=_date_value(start), status=RentalContract.Status.ACTIVE, created_by=actor)
                contracts[key] = contract
            RentalAsset.objects.create(contract=contract, asset=asset, vendor_equipment_number=_text(legacy), rate_amount=_decimal(cost), expected_return_on=_date_value(return_date))
        prior_asset = asset
        messages = ["Placeholder serial preserved only in source staging"] if placeholder else None
        register(ws, row, asset, needs_review, messages)

    # iPads: only rows carrying real serials are assets; all allocation slots remain staged.
    ws = workbook["IT-IPADs"]
    category = _category("Tablets", "TABLET", True)
    for row in range(2, ws.max_row + 1):
        if not _has_values(ws, row):
            continue
        slot, hidden_asset, apparent_tag, team, serial, notes, source_status, returned = [ws.cell(row, col).value for col in range(1, 9)]
        if not serial:
            _stage(batch, ws, row, EquipmentImportRow.Status.ARTIFACT, {"slot": _text(slot)}, ["Unallocated template slot preserved"])
            total += 1
            excluded += 1
            continue
        asset = _create_asset(batch=batch, actor=actor, sequence=sequence, sheet=ws.title, row=row, category=category, legacy_tag=apparent_tag, name="Apple iPad", manufacturer="Apple", ownership=Asset.Ownership.OWNED, status=_status(source_status), location=_location(team), notes=notes, attributes={"allocation_slot": _text(slot), "source_return_date": _json_value(returned)}, review=True, review_notes="Source Employee header contains an asset tag; assignee requires confirmation")
        sequence += 1
        _identifier(asset, AssetIdentifier.Kind.SERIAL, serial, "IT-IPADs", True)
        register(ws, row, asset, True, ["Header/value mismatch requires confirmation"])

    # Preserve every other populated source row without creating demo/person assets.
    for sheet_name in ("Black Box Laptop", "Meta Items", "Meta Assets"):
        ws = workbook[sheet_name]
        for row in range(2, ws.max_row + 1):
            if not _has_values(ws, row):
                continue
            _stage(batch, ws, row, EquipmentImportRow.Status.EXCLUDED, messages=["Template, demo, or identity/access row excluded from equipment registry"])
            total += 1
            excluded += 1

    # Historical maintenance event, linked through imported equipment-number identifier.
    ws = workbook["Maintanance Schedule"]
    for row in range(2, ws.max_row + 1):
        if not _has_values(ws, row):
            continue
        event_date, vehicle_ref, maintenance_type, service_location, duration, owner, notes = [ws.cell(row, col).value for col in range(1, 8)]
        identifier = AssetIdentifier.objects.filter(kind=AssetIdentifier.Kind.EQUIPMENT_NUMBER, normalized_value=_text(vehicle_ref).upper()).select_related("asset").first()
        if identifier:
            vendor = None
            if _text(owner):
                vendor, _ = EquipmentVendor.objects.get_or_create(name=_text(owner))
            dt = _date_value(event_date)
            completed_at = timezone.make_aware(datetime.combine(dt, time.min)) if dt else timezone.now()
            order = MaintenanceWorkOrder.objects.create(work_order_number=f"EQ-WO-IMP-{row:04d}", asset=identifier.asset, title=_text(maintenance_type) or "Imported maintenance", problem_description=_text(notes), status=MaintenanceWorkOrder.Status.COMPLETED, out_of_service=False, completed_at=completed_at, work_performed=f"Imported maintenance at {_text(service_location)}; duration {_text(duration)}", vendor=vendor, opened_by=actor, completed_by=actor)
            _stage(batch, ws, row, EquipmentImportRow.Status.ACCEPTED, {"work_order": order.work_order_number}, asset=identifier.asset)
            accepted += 1
        else:
            _stage(batch, ws, row, EquipmentImportRow.Status.REVIEW, messages=["Vehicle reference could not be linked"])
            review += 1
        total += 1

    batch.status = EquipmentImportBatch.Status.COMPLETED
    batch.total_rows = total
    batch.accepted_rows = accepted
    batch.review_rows = review
    batch.excluded_rows = excluded
    batch.completed_at = timezone.now()
    batch.summary = {
        "assets_created": len(created_assets),
        "asset_quantity_represented": sum(asset.quantity for asset in created_assets),
        "sheets": workbook.sheetnames,
    }
    batch.save(update_fields=("status", "total_rows", "accepted_rows", "review_rows", "excluded_rows", "completed_at", "summary", "updated_at"))
    return batch, True
